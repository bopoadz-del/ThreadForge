"""M1 B05–B09: full TrainingTestCases vendor set, XSD, coverage, branches, xlsx."""

from __future__ import annotations

from pathlib import Path

from threadforge.dexpi_public import (
    entity_counts,
    file_sha256,
    load_pins,
    vendor_xmls,
)
from threadforge.ingest_dexpi import parse_dexpi_xml

ROOT = Path(__file__).resolve().parents[1]
DEXPI13 = ROOT / "fixtures" / "public" / "dexpi13"


def test_b05_vendor_xmls_ingest_and_match_pins():
    xmls = vendor_xmls()
    pins = load_pins()["counts"]
    assert len(xmls) == len(pins) == 35
    exceptions: list[str] = []
    mismatches: list[str] = []
    for path in xmls:
        try:
            graph = parse_dexpi_xml(path)
        except Exception as exc:  # noqa: BLE001
            exceptions.append(f"{path.name}:{type(exc).__name__}:{exc}")
            continue
        got = entity_counts(graph)
        exp = pins[path.name]
        assert file_sha256(path) == exp["sha256"]
        assert path.stat().st_size == exp["bytes"]
        for key in ("pipelines", "equipment", "nozzles", "instruments"):
            if got[key] != exp[key]:
                mismatches.append(f"{path.name}:{key}={got[key]} want={exp[key]}")
    assert exceptions == []
    assert mismatches == []


def test_b05_manifest_sha_matches_files():
    import json

    manifest = json.loads((DEXPI13 / "fetch_manifest.json").read_text(encoding="utf-8"))
    by_name = {row["name"]: row for row in manifest if row.get("sha256") and str(row.get("name", "")).endswith(".xml")}
    pins = load_pins()["counts"]
    for name, exp in pins.items():
        row = by_name[name]
        assert row["sha256"] == exp["sha256"]
        assert row["bytes"] == exp["bytes"]
        assert row.get("license") == "CC-BY-4.0"


def test_b06_each_file_validates_xmlschema_known_deltas_documented():
    from threadforge.dexpi_public import file_known_deltas
    from threadforge.ingest_dexpi import validate_xsd

    for path in vendor_xmls():
        result = validate_xsd(path)
        documented = file_known_deltas(path.name)
        assert result.get("engine") == "xmlschema"
        if documented:
            assert result.get("ok") is False
            assert result.get("known_deltas") == documented
        else:
            assert result.get("ok") is True
            assert result.get("status") == "validated"
            assert result.get("known_deltas") == []
            assert len(result.get("errors") or []) == 0
            assert "4.1.1" in str(result.get("xsd") or "") or result.get("schema_version")


PINNED_VENDOR_ONLY_GAPS = [
    {"name": "AVEVA ComponentClass URI dictionary", "vendor": "AVEVA"},
    {"name": "Hexagon Smart P&ID ComponentClass URI dictionary", "vendor": "Hexagon"},
    {"name": "Autodesk Plant 3D ComponentClass URI dictionary", "vendor": "Autodesk"},
]


def test_b07_coverage_gaps_empty_vendor_only_pinned():
    from threadforge.ingest_dexpi import DEXPI_COVERAGE_GAPS, VENDOR_ONLY_GAPS, coverage_report, load_fixture

    assert DEXPI_COVERAGE_GAPS == []
    assert VENDOR_ONLY_GAPS == PINNED_VENDOR_ONLY_GAPS
    cov = coverage_report()
    assert cov["gaps"] == []
    assert cov["vendor_only_gaps"] == PINNED_VENDOR_ONLY_GAPS
    for name in (
        "PipingComponent subtypes",
        "InstrumentationLoop",
        "SignalLine",
        "ActuatingSystem",
        "InlineComponent",
        "PipeTee",
        "PipeCross",
        "PropertyBreak",
        "SpecBreak",
        "Insulation",
        "Tracing",
    ):
        assert name in cov["supported_elements"]
    g = load_fixture("C01V04-VER.EX01.xml")
    assert any(c.get("component_class") == "PipeTee" for c in g.piping_components.values())
    assert g.instrumentation_loops
    assert g.signal_lines
    assert g.actuating_systems
    assert g.inline_components
    assert isinstance(g.property_breaks, dict)
    assert isinstance(g.spec_breaks, dict)
    g3 = load_fixture("C03V04-VER.EX02.xml")
    assert int(g3.metadata.get("insulation_count") or 0) >= 1
    assert int(g3.metadata.get("tracing_count") or 0) >= 1


def test_b08_c03_branch_count_and_tee_stub_start():
    from threadforge.ingest_dexpi import load_fixture
    from threadforge.routing import fitting_stub_xyz, nozzle_point, route_pipeline

    pins = load_pins()["counts"]
    g3 = load_fixture("C03V04-VER.EX02.xml")
    pinned = int(pins["C03V04-VER.EX02.xml"]["branches"])
    computed = int(g3.metadata.get("branch_count") or 0)
    assert computed == pinned == 0
    assert g3.metadata.get("tee_count") == 0

    g = load_fixture("C01V04-VER.EX01.xml")
    assert int(g.metadata.get("branch_count") or 0) == int(pins["C01V04-VER.EX01.xml"]["branches"])
    assert g.branches
    branch_pipes = [
        p
        for p in g.pipelines.values()
        if p.metadata.get("branch_route") and (p.from_tag in g.branches or p.to_tag in g.branches)
    ]
    assert branch_pipes
    for pipe in branch_pipes:
        route = route_pipeline(g, pipe)
        fitting = pipe.from_tag if pipe.from_tag in g.branches else pipe.to_tag
        stub = fitting_stub_xyz(g, fitting)
        assert stub is not None
        start = (route["points"][0]["x"], route["points"][0]["y"], route["points"][0]["z"])
        assert start == stub
        assert route.get("geometry_source") == "tee_stub"
        assert route.get("start_source") == "tee_stub"
        for nid in g.nozzles:
            npt = nozzle_point(g, nid)
            if npt is None:
                continue
            assert start != npt


def test_b09_c01_xlsx_row_counts_and_columns():
    from openpyxl import load_workbook

    from threadforge.exporters.xlsx import (
        INSTRUMENT_COLUMNS,
        LINE_COLUMNS,
        TIEIN_COLUMNS,
        VALVE_COLUMNS,
        export_lists_xlsx,
        list_row_counts,
    )
    from threadforge.ingest_dexpi import load_fixture

    g = load_fixture("C01V04-VER.EX01.xml")
    pinned = load_pins()["counts"]["C01V04-VER.EX01.xml"]["xlsx"]
    computed = list_row_counts(g)
    assert computed == pinned
    import tempfile

    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "lists.xlsx"
        art = export_lists_xlsx(g, path)
        assert art.status == "ready"
        wb = load_workbook(path)
        expect = {
            "lines": LINE_COLUMNS,
            "valves": VALVE_COLUMNS,
            "instruments": INSTRUMENT_COLUMNS,
            "tie_ins": TIEIN_COLUMNS,
        }
        for sheet, cols in expect.items():
            ws = wb[sheet]
            header = [c.value for c in ws[1]]
            assert header == cols
            data_rows = sum(1 for i, row in enumerate(ws.iter_rows(min_row=2), start=2) if any(c.value not in (None, "") for c in row))
            assert data_rows == pinned[sheet]
        docs = (ROOT / "docs" / "exports.md").read_text(encoding="utf-8")
        for col in LINE_COLUMNS + VALVE_COLUMNS + INSTRUMENT_COLUMNS + TIEIN_COLUMNS:
            assert f"`{col}`" in docs
