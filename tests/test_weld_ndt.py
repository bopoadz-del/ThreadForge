"""B20: weld map + 5% RT from B31.3 341.4.1; csv/xlsx; reconcile B16."""

from __future__ import annotations

import csv

from openpyxl import load_workbook

from threadforge.generators import write_pcf_text
from threadforge.spooling import crafted_straight_30m
from threadforge.tables import B31_3_341_4_1_CITE, B31_3_341_4_1_NORMAL_RT_PCT
from threadforge.weld_ndt import WELD_CSV_COLUMNS, export_weld_ndt, n_rt_required


def test_ndt_pct_cited_b31_3_341_4():
    assert B31_3_341_4_1_NORMAL_RT_PCT == 5.0
    assert "341.4.1" in B31_3_341_4_1_CITE
    assert "5%" in B31_3_341_4_1_CITE or "5 %" in B31_3_341_4_1_CITE
    assert n_rt_required(0) == 0
    assert n_rt_required(1) == 1
    assert n_rt_required(2) == 1  # ceil(0.10) = 1 ≥ 5%
    assert n_rt_required(40) == 2


def test_weld_map_reconciles_b16_and_files(tmp_path):
    g = crafted_straight_30m()
    write_pcf_text(g, "LINE-CRAFT-30M")
    b16 = int(g.routes["LINE-CRAFT-30M"]["spools"]["weld_count"])
    art = export_weld_ndt(g, tmp_path)
    payload = art.payload
    assert payload["weld_count"] == b16 == 2
    assert payload["b16_weld_count"] == b16
    assert payload["reconcile"] is True
    assert payload["rt_pct"] == 5.0
    assert "341.4.1" in payload["citation"]
    assert payload["rt_selected"] == n_rt_required(2)
    rows = payload["rows"]
    assert all(r["vt_pct"] == 100.0 for r in rows)
    assert all(r["weld_type"] == "BW" for r in rows)
    assert all(abs(float(r["wall_mm"]) - 7.11) < 0.05 for r in rows)  # B36.10 6" Sch40
    csv_path = tmp_path / "weld" / "weld_ndt.csv"
    xlsx_path = tmp_path / "weld" / "weld_ndt.xlsx"
    with csv_path.open(encoding="utf-8") as fh:
        got = list(csv.DictReader(fh))
    assert [r["weld_id"] for r in got] == [r["weld_id"] for r in rows]
    assert list(got[0].keys()) == WELD_CSV_COLUMNS
    wb = load_workbook(xlsx_path)
    ws = wb["weld_ndt"]
    assert [c.value for c in ws[1]] == WELD_CSV_COLUMNS
    assert ws.max_row - 1 == b16
